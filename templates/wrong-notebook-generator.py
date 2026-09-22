"""
生成带格式的错题本 Excel。
用法：
    pip install openpyxl
    python wrong-notebook-generator.py
输出：
    错题本.xlsx（含错题记录、复习计划、统计看板三个工作表）
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

# ---------- 样式 ----------
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="4472C4")
sub_fill = PatternFill("solid", fgColor="D9E2F3")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

def style_body(ws, start_row, end_row, ncols):
    for r in range(start_row, end_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = wrap
            cell.border = border

# ---------- Sheet 1: 错题记录 ----------
ws1 = wb.active
ws1.title = "错题记录"
headers1 = [
    "编号", "日期", "年级", "科目", "教材版本", "章节/知识点",
    "题目摘要", "我的错误", "错因分类", "错因细化",
    "正确思路", "关键步骤", "易错点提醒",
    "变式题", "变式答案", "掌握标记", "复习次数", "备注"
]
ws1.append(headers1)
style_header(ws1, 1, len(headers1))

sample = [
    "001", "2026-09-22", "高一", "数学", "人教版",
    "必修一 函数单调性", "f(x)=x²-2ax+3在[1,2]递增，求a",
    "得 a≤1.5", "概念", "忽略对称轴与区间左端点的关系，并把区间中点当成边界",
    "开口向上，递增区间在对称轴右侧", "对称轴 x=a，左端点为1，需 a≤1",
    "先画图，再比较对称轴和区间左端点，并检查等号", "f(x)=x²-2ax+3 在[0,3]递增，求a",
    "a≤0", "未掌握", "1", "注意等号"
]
ws1.append(sample)
style_body(ws1, 2, 2, len(headers1))

widths1 = [6, 12, 8, 8, 10, 18, 22, 18, 10, 18, 22, 22, 18, 22, 12, 10, 8, 12]
for i, w in enumerate(widths1, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w
ws1.freeze_panes = "A2"

# ---------- Sheet 2: 复习计划 ----------
ws2 = wb.create_sheet("复习计划")
headers2 = ["编号", "题目摘要", "首次记录", "第1天", "第3天", "第7天", "第15天",
            "第1次结果", "第2次结果", "第3次结果", "第4次结果", "最终状态"]
ws2.append(headers2)
style_header(ws2, 1, len(headers2))
ws2.append(["001", "二次函数递增求参数", "2026-09-22", "2026-09-23",
            "2026-09-25", "2026-09-29", "2026-10-07",
            "未掌握", "", "", "", "复习中"])
style_body(ws2, 2, 2, len(headers2))
widths2 = [6, 24, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12]
for i, w in enumerate(widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"

# ---------- Sheet 3: 统计看板 ----------
ws3 = wb.create_sheet("统计看板")
stats = [
    ["指标", "数值", "说明"],
    ["总错题数", "=COUNTA(错题记录!A2:A1000)", "错题记录表自动统计"],
    ["未掌握", '=COUNTIF(错题记录!P2:P1000,"未掌握")', "需重点复习"],
    ["模糊", '=COUNTIF(错题记录!P2:P1000,"模糊")', "需巩固"],
    ["已掌握", '=COUNTIF(错题记录!P2:P1000,"已掌握")', "可降低复习频率"],
    ["审题错误", '=COUNTIF(错题记录!I2:I1000,"审题")', "错因分布"],
    ["概念错误", '=COUNTIF(错题记录!I2:I1000,"概念")', "错因分布"],
    ["计算错误", '=COUNTIF(错题记录!I2:I1000,"计算")', "错因分布"],
    ["方法错误", '=COUNTIF(错题记录!I2:I1000,"方法")', "错因分布"],
    ["表达错误", '=COUNTIF(错题记录!I2:I1000,"表达")', "错因分布"],
    ["心态错误", '=COUNTIF(错题记录!I2:I1000,"心态")', "错因分布"],
    ["掌握率", '=IF(B2=0,0,B5/B2)', "已掌握/总数"],
]
for row in stats:
    ws3.append(row)
style_header(ws3, 1, 3)
style_body(ws3, 2, len(stats), 3)
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 32
ws3.column_dimensions["C"].width = 24

wb.save("错题本.xlsx")
print("已生成：错题本.xlsx")
